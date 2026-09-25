#!/usr/bin/env python3
"""Build the Excel review sheet a student fills in, one row per event × camera.

Only forecast exceedance events are listed (the HTF periods the iOS app reported);
--include-controls adds below-threshold control events.

For every event folder with captured frames, each camera gets a row with the HTF
period (the forecast hours at/above threshold, local time and UTC), the HTF point,
the camera and its distance, a link to its image folder, and a yellow
"Has the image flooded?" cell (drop-down: Yes / No / NaN / Invalid) plus optional notes.

Re-run it whenever new events arrive: rows are regenerated, but answers already
typed (matched by event, source and camera) are kept. Whether an event is a forecast
exceedance or a control is kept in a hidden column so the reviewer is not biased.

Usage:
    python3 tools/make_review_sheet.py --events "<Box folder>" [--out "<Box folder>/HTF_camera_review.xlsx"]
"""

import argparse
import json
import re
from datetime import datetime, timezone
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

FONT = "Arial"
YELLOW = PatternFill("solid", start_color="FFFF00")
HEADER_FILL = PatternFill("solid", start_color="1F4E78")
THIN = Side(style="thin", color="BFBFBF")
ANSWERS = ["Yes", "No", "NaN", "Invalid"]   # only Yes/No are scored; NaN and Invalid are excluded
OLD_ANSWERS = {"Y": "Yes", "N": "No"}          # answers typed in earlier versions of the sheet
SOURCE_NAME = {"traffic": "Traffic camera (state DOT)", "usgs": "USGS river/coast camera",
               "windy": "Windy webcam", "webcoos": "WebCOOS coastal camera"}

# (header, width, key). Order = column order on the Review sheet.
COLUMNS = [
    ("Event ID", 34, "event_id"),
    ("HTF ID", 8, "htf_id"),
    ("HTF latitude", 11, "lat"),
    ("HTF longitude", 12, "lon"),
    ("Time zone", 18, "tz"),
    ("HTF period start (local)", 20, "start_local"),
    ("HTF period end (local)", 20, "end_local"),
    ("HTF period start (UTC)", 19, "start_utc"),
    ("HTF period end (UTC)", 19, "end_utc"),
    ("Camera source", 24, "source"),
    ("Camera ID", 26, "camera_id"),
    ("Camera name", 38, "camera_name"),
    ("Distance to HTF point (km)", 13, "distance_km"),
    ("Number of images", 10, "n_images"),
    ("First image (local)", 20, "first_local"),
    ("Last image (local)", 20, "last_local"),
    ("Old images to ignore", 11, "n_stale"),
    ("Image folder", 30, "folder"),
    ("Location map", 12, "map"),
    ("Has the image flooded?", 14, "answer"),
    ("Notes (optional)", 40, "notes"),
    ("Event type (hidden)", 14, "kind"),
]
COL = {key: i + 1 for i, (_, _, key) in enumerate(COLUMNS)}


def local(ts, tz):
    """'2026-09-24T23:15:00Z' → '2026-09-24 19:15' in tz (or UTC if tz unknown)."""
    if not ts:
        return ""
    dt = datetime.fromisoformat(ts.replace("Z", "+00:00"))
    if tz:
        try:
            dt = dt.astimezone(ZoneInfo(tz))
        except Exception:
            pass
    return dt.strftime("%Y-%m-%d %H:%M")


def utc(ts):
    return datetime.fromisoformat(ts.replace("Z", "+00:00")).astimezone(timezone.utc).strftime("%Y-%m-%d %H:%M") if ts else ""


def safe(name):
    return re.sub(r"[^A-Za-z0-9._-]+", "_", str(name)).strip("_")[:80]


def collect_rows(events_dir):
    rows = []
    for path in sorted(Path(events_dir).glob("*/*/event.json")):
        ev = json.loads(path.read_text())
        tz = ev.get("time_zone")
        w = ev["window"]
        by_folder = {f"{c['source']}_{safe(c['id'])}": c for c in ev["cameras"]}
        caps_by_cam = {}
        for cap in ev["captures"]:
            caps_by_cam.setdefault((cap["source"], str(cap["camera_id"])), []).append(cap)
        # One row per camera folder on disk: event.json may not list every frame in it
        # (e.g. frames synced from an earlier run), so images are counted from the files.
        for cam_dir in sorted(d for d in path.parent.iterdir() if d.is_dir()):
            stamps = sorted(f.stem for f in cam_dir.glob("*.jpg"))
            cam = by_folder.get(cam_dir.name)
            if not stamps or not cam:
                continue
            source, cam_id = cam["source"], str(cam["id"])
            caps = caps_by_cam.get((source, cam_id), [])
            times = [re.sub(r"T(\d\d)-(\d\d)Z$", r"T\1:\2:00Z", s) for s in stamps]
            folder = cam_dir.relative_to(events_dir)
            rows.append({
                "event_id": ev["event_id"], "htf_id": ev["htf_id"],
                "lat": round(ev["lat"], 4), "lon": round(ev["lon"], 4), "tz": tz or "UTC",
                "start_local": local(w["start"], tz), "end_local": local(w["end"], tz),
                "start_utc": utc(w["start"]), "end_utc": utc(w["end"]),
                "source": SOURCE_NAME.get(source, source), "source_key": source,
                "camera_id": str(cam_id), "camera_name": cam.get("name", ""),
                "distance_km": cam.get("distance_km"),
                "n_images": len(stamps), "first_local": local(times[0], tz), "last_local": local(times[-1], tz),
                "n_stale": sum(1 for c in caps if c.get("stale")),
                "folder": str(folder), "map": "map.png" if (Path(events_dir) / folder / "map.png").exists() else "",
                "kind": ev["kind"],
            })
    rows.sort(key=lambda r: (r["start_utc"], r["htf_id"], r["distance_km"] if r["distance_km"] is not None else 99))
    return rows


def previous_answers(path):
    """{(event_id, camera source label, camera_id): (answer, notes)} from an existing sheet."""
    if not Path(path).exists():
        return {}
    ws = load_workbook(path)["Review"]
    head = {c.value: i for i, c in enumerate(ws[1])}
    need = ("Event ID", "Camera source", "Camera ID", "Has the image flooded?", "Notes (optional)")
    if not all(h in head for h in need):
        return {}
    out = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[head["Event ID"]]:
            ans = r[head["Has the image flooded?"]]
            out[(r[head["Event ID"]], r[head["Camera source"]], str(r[head["Camera ID"]]))] = (
                OLD_ANSWERS.get(str(ans).strip().upper(), ans) if ans is not None else None, r[head["Notes (optional)"]])
    return out


def style_header(ws, headers, widths):
    for i, (h, w) in enumerate(zip(headers, widths), 1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = HEADER_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 45


def build(events_dir, out_path, include_controls=False):
    old = previous_answers(out_path)
    rows = [r for r in collect_rows(events_dir) if include_controls or r["kind"] == "exceedance"]
    wb = Workbook()

    # ── Instructions ──
    ins = wb.active
    ins.title = "Instructions"
    ins.column_dimensions["A"].width = 26
    ins.column_dimensions["B"].width = 95
    lines = [
        ("HTF camera review", None),
        ("", None),
        ("Your task", "For each row on the Review sheet, open the image folder (click the link) and look at the photos "
                      "taken during the HTF period. Decide whether coastal water has flooded land that is normally dry."),
        ("What to fill in", "Only the YELLOW cells on the Review sheet: 'Has the image flooded?' (choose Yes, No, "
                            "NaN or Invalid from the drop-down) and, if useful, 'Notes (optional)'. Do not edit any other column."),
        ("Yes = flooded", "Water on a road, sidewalk, parking lot or yard; water over a seawall, bulkhead or dock; "
                          "standing water that is not there in the same camera's photos at other times."),
        ("No = not flooded", "Water stays in its normal channel, bay or beach; waves on the sand that do not reach "
                             "roads or buildings; a road that is only wet from rain (no standing or moving water)."),
        ("NaN = cannot be seen", "The photos cannot be judged: images missing or broken, night-time/too dark, fog, "
                                 "rain on the lens, or blurred. Write the reason in Notes."),
        ("Invalid = not usable", "The camera cannot show ground-level flooding, e.g. it is on a bridge or overpass far "
                                 "above the ground, or it points away from the water and land. Write the reason in Notes."),
        ("Times", "Photo file names are UTC (e.g. 2026-09-24T23-15Z.jpg = 23:15 UTC). The Review sheet shows the HTF "
                  "period in local time and UTC. Compare photos before, during and after the period: tidal flooding "
                  "rises and then drains away."),
        ("Old images to ignore", "Number of photos in that folder that were more than 30 min old when captured "
                                 "(frozen or slow cameras). Judge by the other photos."),
        ("One row per camera", "An event can have several cameras. Judge each camera by its own photos; a camera "
                               "that cannot see the flooded area can be N while another is Y."),
        ("Location map", "Each camera folder has map.png showing the HTF point, the 5 km search radius, the camera "
                         "and the distance between them (click 'Open map')."),
        ("", None),
        ("Example row (filled in)", None),
    ]
    for i, (a, b) in enumerate(lines, 1):
        ca = ins.cell(row=i, column=1, value=a)
        ca.font = Font(name=FONT, bold=True, size=14 if i == 1 else 10)
        if b:
            cb = ins.cell(row=i, column=2, value=b)
            cb.font = Font(name=FONT, size=10)
            cb.alignment = Alignment(wrap_text=True, vertical="top")
            ins.row_dimensions[i].height = 42
    ex_row = len(lines) + 1
    example = [("Event ID", "20260924T23Z_HTF1036_exceedance"), ("HTF period (local)", "2026-09-24 19:00 → 22:00 (America/New_York)"),
               ("Camera", "Traffic camera (state DOT) · VA-cam-2853 · 1.3 km"),
               ("Has the image flooded?", "Yes"), ("Notes (optional)", "Water across Shore Dr at 20:15, gone by 21:45")]
    for j, (a, b) in enumerate(example):
        ins.cell(row=ex_row + j, column=1, value=a).font = Font(name=FONT, size=10)
        cb = ins.cell(row=ex_row + j, column=2, value=b)
        cb.font = Font(name=FONT, size=10, italic=True)
        if a in ("Has the image flooded?", "Notes (optional)"):
            cb.fill = YELLOW

    # ── Review ──
    ws = wb.create_sheet("Review")
    style_header(ws, [h for h, _, _ in COLUMNS], [w for _, w, _ in COLUMNS])
    dv = DataValidation(type="list", formula1='"' + ",".join(ANSWERS) + '"', allow_blank=True,
                        error="Choose Yes, No, NaN or Invalid from the list.", errorTitle="Has the image flooded?")
    ws.add_data_validation(dv)
    kept = 0
    for r_i, row in enumerate(rows, 2):
        ans, note = old.get((row["event_id"], row["source"], row["camera_id"]), (None, None))
        kept += ans is not None or note is not None
        row = dict(row, answer=ans, notes=note)
        for _, _, key in COLUMNS:
            c = ws.cell(row=r_i, column=COL[key], value=row.get(key))
            c.font = Font(name=FONT, size=10)
            c.border = Border(bottom=THIN)
            c.alignment = Alignment(vertical="center", wrap_text=key in ("camera_name", "notes"))
        link = ws.cell(row=r_i, column=COL["folder"])
        link.hyperlink = row["folder"]
        link.font = Font(name=FONT, size=10, color="0563C1", underline="single")
        if row["map"]:
            m = ws.cell(row=r_i, column=COL["map"], value="Open map")
            m.hyperlink = f"{row['folder']}/map.png"
            m.font = Font(name=FONT, size=10, color="0563C1", underline="single")
        ws.cell(row=r_i, column=COL["distance_km"]).number_format = "0.0"
        for key in ("answer", "notes"):
            ws.cell(row=r_i, column=COL[key]).fill = YELLOW
        ws.cell(row=r_i, column=COL["answer"]).alignment = Alignment(horizontal="center", vertical="center")
        dv.add(ws.cell(row=r_i, column=COL["answer"]))
    last = max(2, len(rows) + 1)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{last}"
    ws.column_dimensions[get_column_letter(COL["kind"])].hidden = True
    ws.cell(row=1, column=COL["answer"]).comment = Comment("Yes = flooding in the photos during the HTF period; No = none; "
                                                           "NaN = photos cannot be judged (night, missing, fog); Invalid = camera "
                                                           "cannot show ground flooding (e.g. on a bridge).", "HTF review")
    ws.cell(row=1, column=COL["folder"]).comment = Comment("Click to open this camera's photos (the link works when this "
                                                           "file is opened from the Box folder).", "HTF review")

    # ── Progress (formulas) ──
    pr = wb.create_sheet("Progress")
    pr.column_dimensions["A"].width = 34
    pr.column_dimensions["B"].width = 12
    a = get_column_letter(COL["answer"])
    rng = f"Review!${a}$2:${a}${last}"
    ev_rng = f"Review!$A$2:$A${last}"
    stats = [("Rows to review (event × camera)", f'=COUNTA({ev_rng})'),
             ("Yes (flooded)", f'=COUNTIF({rng},"Yes")'),
             ("No (not flooded)", f'=COUNTIF({rng},"No")'),
             ("NaN (cannot be seen)", f'=COUNTIF({rng},"NaN")'),
             ("Invalid (camera not usable)", f'=COUNTIF({rng},"Invalid")'),
             ("Not answered yet", "=B2-B3-B4-B5-B6"),
             ("Share done", "=IF(B2=0,0,(B2-B7)/B2)")]
    pr.cell(row=1, column=1, value="Review progress").font = Font(name=FONT, bold=True, size=14)
    for i, (label, f) in enumerate(stats, 2):
        pr.cell(row=i, column=1, value=label).font = Font(name=FONT, size=10)
        c = pr.cell(row=i, column=2, value=f)
        c.font = Font(name=FONT, size=10)
    pr["B8"].number_format = "0%"

    wb.active = wb.sheetnames.index("Review")
    wb.calculation.fullCalcOnLoad = True       # Excel computes the Progress formulas on open
    wb.save(out_path)
    print(f"{out_path}: {len(rows)} rows ({kept} with earlier answers kept)")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--events", required=True, help="folder with <date>/<event_id>/event.json (the Box folder)")
    ap.add_argument("--out", help="workbook path (default: <events>/HTF_camera_review.xlsx)")
    ap.add_argument("--include-controls", action="store_true",
                    help="also list control (below-threshold) events; default is forecast exceedances only")
    args = ap.parse_args()
    build(Path(args.events), Path(args.out or Path(args.events) / "HTF_camera_review.xlsx"), args.include_controls)


if __name__ == "__main__":
    main()
