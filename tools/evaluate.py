#!/usr/bin/env python3
"""Score the HTF forecast against labeled camera evidence.

Labels come from a CSV kept next to the event folders (e.g. in the Box folder), one
row per event, which the sync never overwrites:

    event_id,flooding_observed,confidence,labeled_by,notes
    20260924T22Z_HTF1071_exceedance,yes,high,EK,water over Shore Dr at 23:15Z

flooding_observed: yes/no (or true/false, 1/0). Rows left blank are skipped. The
"label" block inside event.json is used only for events without a CSV row.

Alternatively --review reads the per-camera Excel sheet from make_review_sheet.py:
an event counts as flooded if any camera is Y, not flooded if its answered cameras
are all N, and is skipped if none is answered.

Forecast = "exceedance" events (predicted flooding) vs "control" events (predicted
no flooding). Observation = flooding_observed. Unlabeled events are skipped, and so
are labeled events whose frames are all stale (image older than STALE_AFTER_MIN when
captured), since their pictures do not show the forecast period.

    forecast \\ observed    flooding      no flooding
    exceedance              hit           false alarm
    control                 miss          correct negative

Usage:
    python3 tools/evaluate.py --events "<Box folder>" --labels "<Box folder>/labels.csv" \
                              [--min-confidence medium] [--csv results.csv]
    python3 tools/evaluate.py --events "<Box folder>" --template "<Box folder>/labels.csv"
"""

import argparse
import csv
import json
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
RANK = {"low": 0, "medium": 1, "high": 2}


def parse_bool(v):
    v = str(v or "").strip().lower()
    return True if v in ("yes", "y", "true", "1") else False if v in ("no", "n", "false", "0") else None


def read_labels(path):
    labels = {}
    if path and Path(path).exists():
        with open(path, newline="") as fh:
            for row in csv.DictReader(fh):
                obs = parse_bool(row.get("flooding_observed"))
                if row.get("event_id") and obs is not None:
                    labels[row["event_id"].strip()] = {
                        "flooding_observed": obs, "confidence": (row.get("confidence") or "low").strip().lower(),
                        "labeled_by": row.get("labeled_by"), "notes": row.get("notes")}
    return labels


def read_review(path):
    """Per-event labels from the review workbook (column 'Has the image flooded?')."""
    from openpyxl import load_workbook
    ws = load_workbook(path, read_only=True, data_only=True)["Review"]
    rows = ws.iter_rows(values_only=True)
    head = {h: i for i, h in enumerate(next(rows))}
    answers = {}
    for r in rows:
        eid, ans = r[head["Event ID"]], str(r[head["Has the image flooded?"]] or "").strip().upper()
        if eid and ans in ("Y", "N"):
            answers.setdefault(eid, []).append(ans)
    return {eid: {"flooding_observed": "Y" in a, "confidence": "medium", "labeled_by": "review sheet",
                  "notes": f"{a.count('Y')} of {len(a)} cameras Y"} for eid, a in answers.items()}


def write_template(events_dir, path):
    """labels.csv with one row per event (existing rows kept), sorted by event_id."""
    existing = {}
    if Path(path).exists():
        with open(path, newline="") as fh:
            existing = {r["event_id"]: r for r in csv.DictReader(fh)}
    fields = ["event_id", "kind", "flooding_observed", "confidence", "labeled_by", "notes"]
    rows = []
    for p in sorted(Path(events_dir).glob("*/*/event.json")):
        ev = json.loads(p.read_text())
        rows.append(existing.get(ev["event_id"]) or {"event_id": ev["event_id"], "kind": ev["kind"]})
    with open(path, "w", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=fields, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)
    print(f"{path}: {len(rows)} events ({len(rows) - len(existing)} new rows)")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--events", default=str(ROOT / "events"), help="folder with <date>/<event_id>/event.json")
    ap.add_argument("--labels", help="labels CSV (event_id, flooding_observed, confidence, labeled_by, notes)")
    ap.add_argument("--review", help="review workbook (HTF_camera_review.xlsx) instead of / as well as --labels")
    ap.add_argument("--template", help="create/extend a labels CSV with a row per event, then exit")
    ap.add_argument("--min-confidence", choices=RANK, default="low")
    ap.add_argument("--csv", help="also write one row per labeled event")
    args = ap.parse_args()
    if args.template:
        write_template(args.events, args.template)
        return
    csv_labels = read_review(args.review) if args.review else {}
    csv_labels.update(read_labels(args.labels))         # labels.csv rows take precedence

    rows, counts = [], {"hit": 0, "false_alarm": 0, "miss": 0, "correct_negative": 0}
    skipped_stale = 0
    for path in sorted(Path(args.events).glob("*/*/event.json")):
        ev = json.loads(path.read_text())
        lab = csv_labels.get(ev["event_id"]) or ev.get("label") or {}
        obs = lab.get("flooding_observed")
        if obs is None or RANK.get(lab.get("confidence") or "low", 0) < RANK[args.min_confidence]:
            continue
        fresh = [c for c in ev["captures"] if not c.get("stale")]
        if not fresh:
            skipped_stale += 1
            continue
        fc = ev["kind"] == "exceedance"
        outcome = ("hit" if obs else "false_alarm") if fc else ("miss" if obs else "correct_negative")
        counts[outcome] += 1
        last = ev["forecasts"][-1] if ev["forecasts"] else {}
        rows.append({"event_id": ev["event_id"], "htf_id": ev["htf_id"], "kind": ev["kind"], "observed": obs,
                     "outcome": outcome, "peak_ft_mhhw": last.get("peak_ft_mhhw"),
                     "threshold_ft_mhhw": ev["threshold_ft_mhhw"], "margin_ft": last.get("margin_ft"),
                     "frames": len(ev["captures"]), "fresh_frames": len(fresh),
                     "stale_frames": len(ev["captures"]) - len(fresh), "confidence": lab.get("confidence")})

    h, f, m, c = counts["hit"], counts["false_alarm"], counts["miss"], counts["correct_negative"]
    n = h + f + m + c
    print(f"Labeled events: {n}" + (f" (+{skipped_stale} skipped: all frames stale)" if skipped_stale else ""))
    print(f"  hits {h} | false alarms {f} | misses {m} | correct negatives {c}")
    if n:
        pod = h / (h + m) if h + m else float("nan")
        far = f / (h + f) if h + f else float("nan")
        csi = h / (h + f + m) if h + f + m else float("nan")
        print(f"  POD (hit rate)          {pod:.2f}")
        print(f"  FAR (false alarm ratio) {far:.2f}")
        print(f"  CSI (threat score)      {csi:.2f}")
        print(f"  Accuracy                {(h + c) / n:.2f}")
        print("  Note: controls are sampled near-threshold points, not a random sample, so"
              " misses are over-represented relative to the whole coast.")
    if args.csv and rows:
        with open(args.csv, "w", newline="") as fh:
            w = csv.DictWriter(fh, fieldnames=list(rows[0]))
            w.writeheader()
            w.writerows(rows)
        print(f"Wrote {len(rows)} rows → {args.csv}")


if __name__ == "__main__":
    main()
